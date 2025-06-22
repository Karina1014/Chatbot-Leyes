from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from sentence_transformers import SentenceTransformer
from qdrant_client import QdrantClient
from qdrant_client.models import VectorParams, Distance, PointStruct
from docx import Document
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
import google.generativeai as genai
import os

# === Cargar configuración de entorno y validar ===
def cargar_config():
    load_dotenv()
    config = {
        "GOOGLE_API_KEY": os.getenv("GOOGLE_API_KEY"),
        "QDRANT_URL": os.getenv("QDRANT_URL"),
        "QDRANT_API_KEY": os.getenv("QDRANT_API_KEY")
    }

    if not all(config.values()):
        raise EnvironmentError("Faltan variables de entorno requeridas.")
    
    return config

config = cargar_config()

# === Inicializar servicios ===
genai.configure(api_key=config["GOOGLE_API_KEY"])
modelo_gemini = genai.GenerativeModel("gemini-1.5-flash")
model_embeddings = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")

qdrant_client = QdrantClient(
    url=config["QDRANT_URL"],
    api_key=config["QDRANT_API_KEY"]
)

# === Constantes de la app ===
UPLOAD_FOLDER = "docs_upload"
COLLECTION_NAME = "documentos_qdrant"
CHUNK_SIZE = 500
MODEL_DIM = 384
EXCEL_PATH = "registro_chat.xlsx"

os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# === FastAPI App ===
app = FastAPI(title="Chatbot Leyes")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# === Utilidades ===
def inicializar_qdrant():
    qdrant_client.recreate_collection(
        collection_name=COLLECTION_NAME,
        vectors_config=VectorParams(size=MODEL_DIM, distance=Distance.COSINE)
    )

def docx_a_chunks(file_path: str, chunk_size: int = CHUNK_SIZE):
    doc = Document(file_path)
    texto = "\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
    chunks = [texto[i:i + chunk_size] for i in range(0, len(texto), chunk_size)]
    vectores = model_embeddings.encode(chunks)
    return chunks, vectores

def construir_prompt(contexto: str, pregunta: str) -> str:
    return f"""
Eres un asistente que responde preguntas usando SOLO la información del contexto proporcionado.

Responde con un lenguaje natural, claro, profesional y bien organizado.

Responde de forma clara, profesional y organizada, SIN usar ningún formato Markdown (como asteriscos, negritas, listas, ni saltos de línea especiales).

Usa texto plano, separando secciones con punto y seguido o punto y aparte solamente. No uses viñetas, ni símbolos especiales. Usa solo texto simple.

Contexto:
\"\"\"{contexto}\"\"\"

Pregunta: {pregunta}

Respuesta completa y bien formateada:
"""

def guardar_en_excel(pregunta: str, respuesta: str, path: str = EXCEL_PATH):
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws[get_column_letter(1) + "1"] = "Pregunta"
        ws[get_column_letter(2) + "1"] = "Respuesta"

    fila = ws.max_row + 1
    ws[f"A{fila}"] = pregunta
    ws[f"B{fila}"] = respuesta
    wb.save(path)

# === Endpoints ===

@app.post("/documento/subir", summary="Subir documento .docx y cargar a Qdrant")
async def subir_documento(file: UploadFile = File(...)):
    if not file.filename.endswith(".docx"):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos .docx")

    ruta = os.path.join(UPLOAD_FOLDER, file.filename)
    with open(ruta, "wb") as f:
        f.write(await file.read())

    chunks, vectores = docx_a_chunks(ruta)
    inicializar_qdrant()

    puntos = [
        PointStruct(id=i, vector=vectores[i].tolist(), payload={"text": chunks[i]})
        for i in range(len(chunks))
    ]
    qdrant_client.upsert(collection_name=COLLECTION_NAME, points=puntos)

    return {
        "estado": "ok",
        "fragmentos_cargados": len(puntos),
        "archivo": file.filename
    }

class ConsultaChat(BaseModel):
    pregunta: str

@app.post("/chat", summary="Consulta al chatbot usando contexto de documentos")
async def consultar_chat(req: ConsultaChat):
    try:
        vector_pregunta = model_embeddings.encode([req.pregunta])[0]
        resultados = qdrant_client.search(
            collection_name=COLLECTION_NAME,
            query_vector=vector_pregunta,
            limit=4
        )

        contexto = "\n\n".join([r.payload["text"] for r in resultados])
        prompt = construir_prompt(contexto, req.pregunta)

        respuesta = modelo_gemini.generate_content(prompt)
        texto_respuesta = respuesta.text.strip()

        # Guardar pregunta y respuesta en Excel
        guardar_en_excel(req.pregunta, texto_respuesta)

        return {"respuesta": texto_respuesta}

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al generar respuesta: {e}")
